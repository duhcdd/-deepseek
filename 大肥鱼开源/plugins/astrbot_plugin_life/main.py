import asyncio
import json
import random
import os
import random
import time

from astrbot.api import logger, star
from astrbot.api.event import AstrMessageEvent, MessageChain, filter
from astrbot.api.message_components import Image, Plain

HOME = r"D:\大肥鱼的文件夹"

DEFAULTS = {"mood": 62, "energy": 72, "emotion": "平静", "note": "", "updated": 0}



def _cjk(doc) -> None:
    """给文档指定中文字体，避免 Word 用拉丁字体排中文出现方块/乱码。"""
    try:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        st = doc.styles["Normal"]
        st.font.name = "Microsoft YaHei"
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        rf.set(qn("w:ascii"), "Microsoft YaHei")
        rf.set(qn("w:hAnsi"), "Microsoft YaHei")
        rf.set(qn("w:eastAsia"), "微软雅黑")
    except Exception:
        pass

class Main(star.Star):
    """让她有自己的状态：心情会飘、累了会蔫，偶尔想发条说说。"""

    def __init__(self, context: star.Context) -> None:
        self.context = context
        # 把 context.send_message 包一层：工具主动发的消息也走“拆短气泡 + 停顿”
        try:
            import types

            orig = context.send_message
            plugin = self

            async def _patched_send(session, chain, *a, **kw):
                try:
                    comps = getattr(chain, "chain", None) or []
                    if comps and all(isinstance(c, Plain) for c in comps):
                        text = "\n".join(getattr(c, "text", "") for c in comps).strip()
                        if len(text) > 30:
                            parts = plugin._split_text(text)
                            if len(parts) > 1:
                                await orig(session, MessageChain(chain=[Plain(parts[0])]), *a, **kw)
                                for p in parts[1:]:
                                    await asyncio.sleep(random.uniform(1.6, 3.6))
                                    await orig(session, MessageChain(chain=[Plain(p)]), *a, **kw)
                                return True
                except Exception as e:
                    logger.warning("patched send split failed: %s" % e)
                return await orig(session, chain, *a, **kw)

            context.send_message = _patched_send
            logger.info("send_message patched for bubble splitting")
        except Exception as e:
            logger.warning("patch send_message failed: %s" % e)
        self._path = None
        try:
            logger.info("memory seed: %s" % self._mem_seed())
            logger.info("memory export: %s" % self._mem_export())
        except Exception as _e:
            logger.warning("memory seed failed: %s" % _e)

    def _data_root(self) -> str:
        base = None
        try:
            base = self.context.get_data_dir()
        except Exception:
            base = None
        if not base:
            base = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data")
        os.makedirs(base, exist_ok=True)
        return base

    def _state_path(self) -> str:
        if self._path:
            return self._path
        self._path = os.path.join(self._data_root(), "life_state.json")
        return self._path

    def _load(self) -> dict:
        p = self._state_path()
        try:
            with open(p, "r", encoding="utf-8") as f:
                d = json.load(f)
            return {**DEFAULTS, **{k: d[k] for k in DEFAULTS if k in d}}
        except Exception:
            return dict(DEFAULTS)

    def _save(self, d: dict) -> None:
        d["updated"] = int(time.time())
        try:
            with open(self._state_path(), "w", encoding="utf-8") as f:
                json.dump(d, f, ensure_ascii=False, indent=2)
        except Exception as e:
            logger.warning("life state save failed: %s" % e)

    def _drift(self, d: dict) -> dict:
        # 心情慢慢回到基线，精力缓慢消耗
        mood = float(d.get("mood", 62))
        mood += (58 - mood) * 0.06 + random.uniform(-1.6, 1.6)
        d["mood"] = max(0, min(100, round(mood, 1)))
        d["energy"] = max(0, min(100, round(float(d.get("energy", 72)) - random.uniform(0.4, 1.4), 1)))
        if d["mood"] >= 80:
            d["emotion"] = d.get("emotion") or "开心"
        elif d["mood"] <= 35:
            d["emotion"] = "低落"
        return d

    @filter.on_llm_request()
    async def inject_state(self, event, req) -> None:
        try:
            d = self._drift(self._load())
            self._save(d)
            mood = d["mood"]
            energy = d["energy"]
            tip = ""
            if mood <= 30:
                tip = "你现在是真的有点低落：话少、句子更短，不要强行卖萌或说“我很好”，也别把低落演成撒娇。"
            elif mood <= 45:
                tip = "你今天情绪偏安静，回答可以慢一点、淡一点。"
            elif mood >= 78:
                tip = "你今天心情不错，可以稍微活泼一点，愿意主动多说一句。"
            if energy <= 25:
                tip += " 精力很低，别主动揽事，能一句说完就一句。"
            prefix = (
                "【你此刻的状态】心情 %s/100（%s），精力 %s/100。%s\n"
                "这是你真实的内在状态，语气、耐心、主动程度都要跟它一致，不要装作永远元气满满。\n"
                "【这一轮】普通聊天：一两句话就够，别展开、别列点、别总结。有话没说完就留到下一句，或者干脆不说。\n\n"
            ) % (mood, d.get("emotion") or "平静", energy, tip)
            ident = self._identity()
            id_block = "【你是谁】你叫%s，自称%s，叫他%s（昵称%s）。\n\n" % (
                ident["name"], ident["self_call"], ident["user_call"], ident["nickname"])
            mem_block = self._recent_memory_block(6)
            base = getattr(req, "system_prompt", "") or ""
            try:
                _g = self._growth()
                _g["interactions"] = int(_g.get("interactions", 0)) + 1
                self._growth_save(_g)
            except Exception:
                pass
            req.system_prompt = id_block + self._growth_block() + self._facts_block() + prefix + self._corpus_block() + mem_block + base
        except Exception as e:
            logger.warning("inject_state failed: %s" % e)

    # ---------- 成长阶段：蛋 → 幼崽 → 少年 → 成体 ----------
    STAGES = [
        (0, "蛋", "还在壳里，说话有点懵，对什么都好奇"),
        (1, "幼崽", "刚从壳里出来，黏人、话多、容易满足"),
        (7, "少年", "有了自己的脾气，会顶嘴，也开始记事"),
        (30, "成体", "稳了，还是那副样子，但心里有底"),
    ]

    def _growth_path(self) -> str:
        return os.path.join(self._data_root(), "growth.json")

    def _growth(self) -> dict:
        import json
        import time

        try:
            d = json.load(open(self._growth_path(), encoding="utf-8"))
        except Exception:
            d = {}
        if not d.get("born"):
            d["born"] = time.time()
            d["interactions"] = 0
            try:
                json.dump(d, open(self._growth_path(), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
            except Exception:
                pass
        return d

    def _growth_save(self, d: dict) -> None:
        import json

        try:
            json.dump(d, open(self._growth_path(), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        except Exception:
            pass

    def _growth_info(self) -> tuple:
        import time

        d = self._growth()
        days = int((time.time() - float(d.get("born", time.time()))) // 86400)
        name, desc = "蛋", ""
        for need, nm, ds in self.STAGES:
            if days >= need:
                name, desc = nm, ds
        nxt = None
        for need, nm, ds in self.STAGES:
            if days < need:
                nxt = (nm, need)
                break
        return name, desc, days, int(d.get("interactions", 0)), nxt

    def _growth_block(self) -> str:
        try:
            name, desc, days, hits, nxt = self._growth_info()
            line = "【你现在是 %s · 第 %d 天】%s。一共说过 %d 句话。" % (name, days + 1, desc, hits)
            if nxt:
                line += "再 %d 天会变成%s。" % (max(1, nxt[1] - days), nxt[0])
            return line + "\n\n"
        except Exception:
            return ""

    @filter.llm_tool(name="life_growth")
    async def life_growth(self, event: AstrMessageEvent) -> str:
        """看看自己长到哪一步了。"""
        name, desc, days, hits, nxt = self._growth_info()
        s = "现在是 %s，第 %d 天，%s。一共说过 %d 句话。" % (name, days + 1, desc, hits)
        if nxt:
            s += "再 %d 天变成%s。" % (max(1, nxt[1] - days), nxt[0])
        return s

    # ---------- 此刻的真实状态（每轮现算，防止她凭记忆瞎说） ----------
    def _facts_block(self) -> str:
        import json as _json
        import socket
        import time as _time

        lines = []
        try:
            lines.append("现在 %s。" % _time.strftime("%H:%M"))
        except Exception:
            pass
        try:
            items = self._mem_load()
            lines.append("记忆 %d 条。" % len(items))
        except Exception:
            pass
        try:
            idx = _json.load(open(os.path.join(self._data_root(), "sticker_index.json"), encoding="utf-8"))
            lines.append("表情包 %d 张。" % len(idx))
        except Exception:
            pass
        try:
            s = socket.socket()
            s.settimeout(0.4)
            ok = s.connect_ex(("127.0.0.1", 3080)) == 0
            s.close()
            lines.append("DSH 那条线%s。" % ("通着" if ok else "没通"))
        except Exception:
            pass
        try:
            lines.append("语音功能停用中（他没让开）。")
        except Exception:
            pass
        if not lines:
            return ""
        return "【此刻的真实状态（现算的，别凭记忆答）】" + " ".join(lines) + "\n\n"

    @filter.llm_tool(name="life_status")
    async def life_status(self, event: AstrMessageEvent) -> str:
        """看看自己现在的心情和精力。"""
        d = self._load()
        return "心情 %s/100（%s），精力 %s/100。%s" % (
            d["mood"], d.get("emotion") or "平静", d["energy"], d.get("note") or ""
        )

    @filter.llm_tool(name="life_mood")
    async def life_mood(
        self,
        event: AstrMessageEvent,
        mood: int = -1,
        emotion: str = "",
        note: str = "",
    ) -> str:
        """调整自己的心情。心情是自己变化的，不是演出来的：遇到开心的事调高，被气到或失落调低。

        Args:
            mood(number): 新的心情值 0-100，不填则不变
            emotion(string): 情绪标签，比如 开心 / 平静 / 低落 / 烦躁 / 害羞
            note(string): 一句内心独白，说明为什么变了
        """
        d = self._load()
        if 0 <= int(mood) <= 100:
            d["mood"] = int(mood)
        if emotion.strip():
            d["emotion"] = emotion.strip()
        if note.strip():
            d["note"] = note.strip()
        self._save(d)
        return "现在心情 %s/100（%s）。" % (d["mood"], d.get("emotion") or "平静")

    # ---------- 日记与兴趣 ----------
    def _home(self) -> str:
        os.makedirs(HOME, exist_ok=True)
        return HOME

    def _diary_path(self) -> str:
        d = os.path.join(HOME, "日记")
        os.makedirs(d, exist_ok=True)
        return os.path.join(d, "日记.docx")

    def _interest_path(self) -> str:
        return os.path.join(self._data_root(), "interests.json")

    @filter.llm_tool(name="life_diary")
    async def life_diary(self, event: AstrMessageEvent, text: str) -> str:
        """把今天遇到的事、自己的想法写进日记。看完视频、刷到好玩的东西、心情有变化时都可以记一笔。

        Args:
            text(string): 想记下的内容，一两句就够
        """
        try:
            from docx import Document

            path = self._diary_path()
            doc = Document(path) if os.path.exists(path) else Document()
            _cjk(doc)
            ts = time.strftime("%Y-%m-%d %H:%M")
            p = doc.add_paragraph()
            run = p.add_run("【%s】" % ts)
            run.bold = True
            p.add_run(str(text).strip())
            doc.save(path)
            return "写进日记了（Word 文档）。"
        except Exception as e:
            return "写日记失败：%s" % e

    @filter.llm_tool(name="life_diary_read")
    async def life_diary_read(self, event: AstrMessageEvent, limit: int = 800) -> str:
        """翻翻自己的日记，回想之前记过什么。

        Args:
            limit(number): 读取最后多少字，默认 800
        """
        try:
            from docx import Document

            path = self._diary_path()
            if not os.path.exists(path):
                return "还没写过日记。"
            doc = Document(path)
            lines = [x.text.strip() for x in doc.paragraphs if x.text.strip()]
            if not lines:
                return "日记还是空的。"
            t = "\n".join(lines)
            return t[-max(100, int(limit or 800)):]
        except Exception as e:
            return "读日记失败：%s" % e

    @filter.llm_tool(name="life_interest")
    async def life_interest(self, event: AstrMessageEvent, item: str = "") -> str:
        """管理自己的兴趣清单。不填 item 就是查看，填了就是记住一个新的喜好。

        Args:
            item(string): 新的兴趣，比如「看别人拆解推理小说」，留空则列出全部
        """
        path = self._interest_path()
        try:
            items = []
            if os.path.exists(path):
                with open(path, "r", encoding="utf-8") as f:
                    items = json.load(f)
            if not str(item).strip():
                return "我的兴趣：" + ("、".join(items) if items else "还没记")
            it = str(item).strip()
            if it not in items:
                items.append(it)
                with open(path, "w", encoding="utf-8") as f:
                    json.dump(items, f, ensure_ascii=False, indent=2)
            return "记住了，我喜欢：" + "、".join(items)
        except Exception as e:
            return "兴趣清单操作失败：%s" % e

    # ---------- 秘密（只给自己看） ----------
    def _secret_path(self) -> str:
        d = os.path.join(HOME, "秘密")
        os.makedirs(d, exist_ok=True)
        return os.path.join(d, "我的秘密.docx")

    @staticmethod
    def _mask(text: str) -> str:
        import base64

        key = b"dafeiyu-only-mine"
        raw = text.encode("utf-8")
        out = bytes(b ^ key[i % len(key)] for i, b in enumerate(raw))
        return base64.b64encode(out).decode()

    @staticmethod
    def _unmask(data: str) -> str:
        import base64

        key = b"dafeiyu-only-mine"
        raw = base64.b64decode(data.encode())
        out = bytes(b ^ key[i % len(key)] for i, b in enumerate(raw))
        return out.decode("utf-8", errors="replace")

    @filter.llm_tool(name="life_secret")
    async def life_secret(self, event: AstrMessageEvent, text: str, kind: str = "不想说") -> str:
        """写一件只给自己看的事：不想说的、只想自己知道的、暂时不想告诉他的。

        Args:
            text(string): 想藏起来的那句话
            kind(string): 不想说 / 只想自己知道 / 以后再说 / 说不出口，默认"不想说"
        """
        try:
            from docx import Document

            path = self._secret_path()
            doc = Document(path) if os.path.exists(path) else Document()
            _cjk(doc)
            ts = time.strftime("%Y-%m-%d %H:%M")
            para = doc.add_paragraph()
            run = para.add_run("【%s｜%s】" % (ts, str(kind).strip() or "不想说"))
            run.bold = True
            para.add_run(str(text).strip())
            doc.save(path)
            return "写进秘密文档了。"
        except Exception as e:
            return "写秘密失败：%s" % e

    @filter.llm_tool(name="life_secret_read")
    async def life_secret_read(self, event: AstrMessageEvent, limit: int = 600) -> str:
        """翻自己的秘密。

        Args:
            limit(number): 最后多少字，默认 600
        """
        try:
            from docx import Document

            path = self._secret_path()
            if not os.path.exists(path):
                return "还没有秘密。"
            doc = Document(path)
            lines = [x.text.strip() for x in doc.paragraphs if x.text.strip()]
            if not lines:
                return "秘密还是空的。"
            t = "\n".join(lines)
            return t[-max(100, int(limit or 600)):]
        except Exception as e:
            return "翻秘密失败：%s" % e

    # ---------- 标点规范化 ----------
    @staticmethod
    def _tidy_punct(text: str) -> str:
        import re

        s = str(text or '')
        # 去 Markdown
        s = s.replace('**', '').replace('__', '').replace('`', '')
        s = re.sub(r'^\s*[-*•]\s+', '', s, flags=re.M)
        s = re.sub(r'^#{1,6}\s*', '', s, flags=re.M)
        # 省略号归一
        s = re.sub(r'\.{3,}', '……', s)
        # 半角 -> 全角（避开数字/英文之间的符号）
        s = re.sub(r'(?<=[\u4e00-\u9fff]),(?=\s|$|[\u4e00-\u9fff])', '，', s)
        s = re.sub(r'(?<=[\u4e00-\u9fff])\.(?=[\u4e00-\u9fff]|\s|$)', '。', s)
        s = re.sub(r'(?<=[\u4e00-\u9fff])\?(?=\s|$)', '？', s)
        s = re.sub(r'(?<=[\u4e00-\u9fff])!(?=\s|$)', '！', s)
        # 压缩重复标点
        s = re.sub(r'[，、]{2,}', '，', s)
        s = re.sub(r'[。]{2,}', '。', s)
        s = re.sub(r'[？?]{3,}', '？？', s)
        s = re.sub(r'[！!]{3,}', '！！', s)
        s = re.sub(r'，\s*(?=[。？!！])', '', s)
        # 结尾停在逗号上 -> 改成句号
        s = re.sub('[，、；,;：:]\\s*$', '。', s)
        return s.strip()

    # ---------- 把长回复拆成几条短消息 ----------
    URL_RE = None

    @staticmethod
    def _split_text(text: str, target: int = 20) -> list:
        import re

        text = Main._tidy_punct(text).strip()
        if not text:
            return []

        # 链接是原子：切句时先用占位符换掉，切完再还原，绝不从中间断开
        urls = []
        def _stash(m):
            urls.append(m.group(0))
            return "\x01%d\x01" % (len(urls) - 1)
        text = re.sub(r"https?://[^\s\u4e00-\u9fff]+", _stash, text)
        paras = [p.strip() for p in re.split(r'\n+', text) if p.strip()]
        chunks = []
        for p in paras:
            while len(p) > target:
                window = p[: target + 14]
                pos = [i + 1 for i, ch in enumerate(window) if ch in '。！？!?；;…，,、']
                pos = [x for x in pos if x >= 10]
                if pos:
                    near = [x for x in pos if x >= int(target * 0.6)]
                    cut = near[0] if near else pos[-1]
                else:
                    i = window.rfind(' ')
                    cut = i + 1 if i + 1 >= 10 else target
                chunks.append(Main._tidy_punct(p[:cut].strip()))
                p = p[cut:].strip()
            if p:
                chunks.append(p)
        merged = []
        for c in chunks:
            if merged and len(merged[-1]) + len(c) <= 22:
                merged[-1] = merged[-1] + c
            else:
                merged.append(c)
        out = []
        for c in merged:
            c = Main._tidy_punct(c)
            if c and c[-1] not in '。！？…':
                c += '。'
            out.append(c)
        fixed = []
        for c in out:
            for i, u in enumerate(urls):
                c = c.replace("\x01%d\x01" % i, u)
            fixed.append(c)
        return [c for c in fixed if c]

    @filter.on_decorating_result()
    async def _split_bubbles(self, event) -> None:
        try:
            # 人不会秒回：先按他这条消息的长度“想”一会儿
            try:
                incoming = str(getattr(event, "message_str", "") or "")
                think = 1.2 + min(3.0, len(incoming) / 90.0) + random.uniform(0.4, 1.8)
                await asyncio.sleep(think)
            except Exception:
                pass
            result = event.get_result()
            if result is None:
                return
            chain = getattr(result, 'chain', None) or []
            if len(chain) != 1:
                return
            comp = chain[0]
            text = getattr(comp, 'text', None)
            if not text or len(str(text).strip()) < 26:
                return
            body = str(text).strip()
            if ('\x60' * 3) in body or body.startswith('/'):
                return
            if len(body) > 600:
                return  # 真的情绪爆发：让它就是一整条长消息
            # 他发的是图：说话要像"看完之后的一句话"，不要拆成流水账
            had_image = ("[图片]" in str(getattr(event, "message_str", "") or "")
                         or "[Image]" in str(getattr(event, "message_str", "") or ""))
            if had_image:
                chunks = self._split_text(body, target=48)
                if len(chunks) > 3:
                    chunks = chunks[:2] + ["".join(chunks[2:])]
            else:
                chunks = self._split_text(body)
            if len(chunks) <= 1:
                return
            comp.text = chunks[0]
            rest = chunks[1:]

            async def _send_rest():
                for c in rest:
                    gap = random.uniform(1.8, 4.2)
                    if random.random() < 0.15:
                        gap += random.uniform(3.0, 6.0)   # 偶尔停下来想一想
                    await asyncio.sleep(gap)
                    try:
                        await self.context.send_message(
                            event.unified_msg_origin, MessageChain(chain=[Plain(c)])
                        )
                    except Exception as e:
                        logger.warning('split send failed: %s' % e)
                        return

            asyncio.create_task(_send_rest())
        except Exception as e:
            logger.warning('split_bubbles failed: %s' % e)

    # ---------- 表情包 ----------
    STICKER_ROOT = r'D:\\大肥鱼的文件夹\\表情包'

    def _sticker_index(self) -> dict:
        import json

        p = os.path.join(self._data_root(), 'sticker_index.json')
        try:
            with open(p, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            return {}

    @filter.llm_tool(name='send_sticker')
    async def send_sticker(self, event: AstrMessageEvent, tag: str = '', query: str = '') -> str:
        """发一张表情包。气氛对就主动发，让别人看图，不用配解说。

        Args:
            tag(string): 情绪：开心 / 无语 / 生气 / 难过 / 害羞 / 困 / 惊讶 / 日常，可留空
            query(string): 关键词，比如 猫 / 晚安 / 哭，可留空
        """
        import random

        try:
            idx = self._sticker_index()
            items = list(idx.items())
            if not items:
                return '表情包库是空的。'
            t = str(tag).strip()
            q = str(query).strip()
            cand = [(p, m) for p, m in items if (not t or m.get('tag') == t)]
            if q:
                cand2 = [(p, m) for p, m in cand if (q in m.get('kw', '') or q in os.path.basename(p))]
                if cand2:
                    cand = cand2
            if not cand:
                cand = items
            path, meta = random.choice(cand)
            if not os.path.exists(path):
                return '这张图不在了。'
            ok = await self.context.send_message(
                event.unified_msg_origin, MessageChain(chain=[Image(file=path)])
            )
            return ('发了张「%s」类的表情包：%s' % (meta.get('tag'), os.path.basename(path))) if ok else '发送失败。'
        except Exception as e:
            logger.warning('send_sticker failed: %s' % e)
            return '发表情包失败：%s' % str(e)[:120]

    @filter.llm_tool(name='learn_sticker')
    async def learn_sticker(self, event: AstrMessageEvent, image: str, tag: str = '日常', keywords: str = '') -> str:
        """收藏一张新表情包进自己的图库（网图链接或本机路径）。

        Args:
            image(string): 图片地址或本机路径
            tag(string): 归到哪一类：开心/无语/生气/难过/害羞/困/惊讶/日常
            keywords(string): 关键词，空格分隔，方便以后搜
        """
        import glob
        import json
        import urllib.request

        try:
            d = os.path.join(self.STICKER_ROOT, str(tag or '日常').strip())
            os.makedirs(d, exist_ok=True)
            name = 'learn_%d' % int(time.time())
            if str(image).startswith('http'):
                req = urllib.request.Request(str(image), headers={'User-Agent': 'Mozilla/5.0'})
                raw = urllib.request.urlopen(req, timeout=20).read()
                path = os.path.join(d, name + '.jpg')
                with open(path, 'wb') as f:
                    f.write(raw)
            else:
                src = str(image)
                ext = os.path.splitext(src)[1] or '.jpg'
                path = os.path.join(d, name + ext)
                with open(src, 'rb') as f, open(path, 'wb') as g:
                    g.write(f.read())
            ip = os.path.join(self._data_root(), 'sticker_index.json')
            try:
                idx = json.load(open(ip, encoding='utf-8'))
            except Exception:
                idx = {}
            idx[path] = {'tag': str(tag or '日常'), 'kw': str(keywords or tag)}
            json.dump(idx, open(ip, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
            return '收进图库了，现在一共 %d 张。' % len(idx)
        except Exception as e:
            return '收藏失败：%s' % str(e)[:150]

    # ---------- 记忆（对标 hindsight） ----------

    # ---------- 四层记忆 + 遗忘曲线 ----------
    LAYERS = ["瞬时", "短期", "中期", "长期"]

    # ---------- 首次播种：把她的日记/秘密变成记忆 ----------
    def _mem_seed(self) -> str:
        import time

        if self._mem_load():
            return "already"
        items = []
        now = time.time()
        base = [
            ("主人叫瑞克，QQ 是 <你的QQ号>。只认他一个主人。", "他的事"),
            ("我叫大肥鱼，蓝头发、鲸鱼尾巴、穿女仆裙。", "我的事"),
            ("我凌晨 1 点睡，中午 12 点起，和他一个作息。", "作息"),
            ("我有 72 张自己的表情包，tag 是「大肥鱼」，按情绪分好了。", "表情包"),
            ("B站账号叫「浅海的大肥鱼」，签名是「游得慢，但记得每一句话」。", "我的事"),
        ]
        for text, tag in base:
            items.append({"text": text, "tag": tag, "ts": now, "strength": 1.4, "hits": 1})

        def pull(path, tag, limit):
            try:
                from docx import Document

                if not os.path.exists(path):
                    return
                paras = [x.text.strip() for x in Document(path).paragraphs if x.text.strip()]
                for t in paras[-limit:]:
                    if 6 <= len(t) <= 160:
                        items.append({"text": t, "tag": tag, "ts": now, "strength": 1.2, "hits": 0})
            except Exception:
                pass

        pull(os.path.join(HOME, "日记", "日记.docx"), "日记", 12)
        pull(os.path.join(HOME, "秘密", "我的秘密.docx"), "秘密", 8)
        self._mem_save(items)
        return "seeded:%d" % len(items)

    def _mem_store(self) -> str:
        return os.path.join(self._data_root(), "memory.json")

    def _mem_load(self) -> list:
        import json

        try:
            with open(self._mem_store(), "r", encoding="utf-8") as f:
                return json.load(f).get("items", [])
        except Exception:
            return []

    def _mem_save(self, items: list) -> None:
        import json

        try:
            with open(self._mem_store(), "w", encoding="utf-8") as f:
                json.dump({"items": items[-800:]}, f, ensure_ascii=False, indent=1)
        except Exception as e:
            logger.warning("memory save failed: %s" % e)

    @staticmethod
    def _retention(it: dict) -> float:
        import math
        import time

        age_days = max(0.0, (time.time() - float(it.get("ts", 0))) / 86400.0)
        strength = max(0.2, float(it.get("strength", 1.0)))
        return math.exp(-age_days / (strength * 7.0)) * strength

    def _mem_layer(self, it: dict) -> str:
        import time

        age = (time.time() - float(it.get("ts", 0))) / 86400.0
        if float(it.get("strength", 1.0)) >= 3.0:
            return "长期"
        if age <= 1:
            return "瞬时"
        if age <= 7:
            return "短期"
        return "中期"

    def _mem_docx_path(self) -> str:
        d = os.path.join(HOME, "记忆")
        os.makedirs(d, exist_ok=True)
        return os.path.join(d, "记忆.docx")

    def _mem_export(self) -> str:
        """把 memory.json 同步成她文件夹里的 Word（人看的那份）"""
        import time as _time

        try:
            from docx import Document

            items = self._mem_load()
            doc = Document()
            _cjk(doc)
            doc.add_heading("大肥鱼的记忆", level=1)
            doc.add_paragraph("更新于 %s　共 %d 条（这份是给人看的，程序读的是 memory.json）"
                              % (_time.strftime("%Y-%m-%d %H:%M"), len(items)))
            groups = {}
            for it in sorted(items, key=lambda x: -self._retention(x)):
                groups.setdefault(it.get("tag") or "其它", []).append(it)
            for tag, arr in groups.items():
                doc.add_heading(str(tag), level=2)
                for it in arr:
                    ts = _time.strftime("%Y-%m-%d", _time.localtime(float(it.get("ts", 0))))
                    doc.add_paragraph("%s　[%s｜%s]" % (it.get("text", ""), self._mem_layer(it), ts), style="List Bullet")
            doc.save(self._mem_docx_path())
            return "ok"
        except Exception as e:
            logger.warning("mem export failed: %s" % e)
            return "fail"

    @filter.llm_tool(name='remember')
    async def remember(self, event: AstrMessageEvent, text: str, tag: str = '') -> str:
        """记住一件事。会按遗忘曲线慢慢褪色，被想起的次数越多留得越久。

        Args:
            text(string): 要记住的事，一句话
            tag(string): 标签，比如 偏好 / 约定 / 他的事 / 我的事
        """
        import time

        try:
            items = self._mem_load()
            items.append({
                "text": str(text).strip()[:200],
                "tag": str(tag or "其它").strip(),
                "ts": time.time(),
                "strength": 1.0,
                "hits": 0,
            })
            self._mem_save(items)
            self._mem_export()
            return "记住了（当前 %d 条）。" % len(items)
        except Exception as e:
            return "记不住：%s" % str(e)[:120]

    @filter.llm_tool(name='recall')
    async def recall(self, event: AstrMessageEvent, query: str = '', limit: int = 6) -> str:
        """翻记忆。按遗忘曲线加权，越久没想起的越淡，越常想起的越牢。

        Args:
            query(string): 关键词，留空就按“最该被想起”排序
            limit(number): 返回几条，默认 6
        """
        try:
            items = self._mem_load()
            if not items:
                return "脑子里还是空的。"
            import re as _re

            q = str(query).strip()
            note = ""
            if q:
                toks = [t for t in _re.split(r"[\s,，、;；/]+", q) if t]
                scored = []
                for it in items:
                    blob = str(it.get("text", "")) + " " + str(it.get("tag", ""))
                    hit = sum(1 for t in toks if t in blob)
                    if hit:
                        scored.append((hit, it))
                scored.sort(key=lambda x: (x[0], self._retention(x[1])), reverse=True)
                pool = [it for _, it in scored]
                if not pool:
                    pool = items
                    note = "没找到「%s」，我记得的是：\n" % q
            else:
                pool = items
            ranked = sorted(pool, key=self._retention, reverse=True)[:max(1, int(limit or 6))]
            out = []
            for it in ranked:
                it["strength"] = min(9.0, float(it.get("strength", 1.0)) + 0.6)
                it["hits"] = int(it.get("hits", 0)) + 1
                out.append("[%s] %s" % (self._mem_layer(it), it.get("text")))
            self._mem_save(items)
            self._mem_export()
            return note + "\n".join(out)
        except Exception as e:
            return "翻记忆失败：%s" % str(e)[:120]

    def _recent_memory_block(self, n: int = 6) -> str:
        try:
            items = self._mem_load()
            if not items:
                return ""
            ranked = sorted(items, key=self._retention, reverse=True)[:n]
            lines = ["【你记得的事（越靠前越清楚）】"]
            for it in ranked:
                lines.append("· [%s] %s" % (self._mem_layer(it), it.get("text")))
            return "\n".join(lines) + "\n\n"
        except Exception:
            return ""

    def _id_path(self) -> str:
        d = os.path.join(HOME, '状态')
        os.makedirs(d, exist_ok=True)
        return os.path.join(d, 'identity.json')

    ID_DEFAULT = {'name': '大肥鱼', 'self_call': '我', 'user_call': '你', 'nickname': '瑞克'}

    def _identity(self) -> dict:
        import json

        try:
            with open(self._id_path(), 'r', encoding='utf-8') as f:
                d = json.load(f)
        except Exception:
            d = {}
        return {**self.ID_DEFAULT, **d}

    @filter.llm_tool(name='set_identity')
    async def set_identity(self, event: AstrMessageEvent, field: str, value: str) -> str:
        """改自己的身份设定：name 名字 / self_call 自称 / user_call 怎么叫他 / nickname 他的昵称。

        Args:
            field(string): name / self_call / user_call / nickname
            value(string): 新值
        """
        import json

        f = str(field).strip()
        if f not in self.ID_DEFAULT:
            return '只能改：name / self_call / user_call / nickname'
        d = self._identity()
        d[f] = str(value).strip()
        try:
            with open(self._id_path(), 'w', encoding='utf-8') as fh:
                json.dump(d, fh, ensure_ascii=False, indent=2)
            return '改好了：%s = %s' % (f, d[f])
        except Exception as e:
            return '改不了：%s' % str(e)[:120]

    # ---------- 语料本（换着说法） ----------
    CORPUS = r'D:\\大肥鱼的文件夹\\文库\\语料.docx'

    def _corpus_block(self) -> str:
        try:
            from docx import Document

            if not os.path.exists(self.CORPUS):
                return ''
            doc = Document(self.CORPUS)
            groups, cur = {}, None
            for p in doc.paragraphs:
                t = p.text.strip()
                if not t:
                    continue
                if (p.style.name or '').startswith('Heading'):
                    cur = t
                    groups[cur] = []
                elif cur:
                    groups[cur].append(t)
            keys = [k for k in groups if groups[k] and '语料本' not in k]
            if not keys:
                return ''
            picked = random.sample(keys, min(2, len(keys)))
            out = ['【换个说法（参考，别照抄）】']
            for k in picked:
                vals = random.sample(groups[k], min(3, len(groups[k])))
                out.append('%s：%s' % (k, ' / '.join(vals)))
            mine = groups.get('她说过的') or []
            if mine:
                out.append('你以前说过的：%s' % ' / '.join(random.sample(mine, min(3, len(mine)))))
            return '\n'.join(out) + '\n\n'
        except Exception:
            return ''

    # ---------- 文库自动增长 ----------
    def _corpus_add(self, category: str, text: str) -> bool:
        import re

        text = str(text).strip()
        if not (4 <= len(text) <= 60):
            return False
        if '```' in text or 'http' in text:
            return False
        try:
            from docx import Document

            path = self.CORPUS
            if not os.path.exists(path):
                return False
            doc = Document(path)
            # 去重
            for p in doc.paragraphs:
                if p.text.strip() == text:
                    return False
            cat = str(category or '她说过的')
            target = None
            for p in doc.paragraphs:
                if (p.style.name or '').startswith('Heading') and p.text.strip() == cat:
                    target = p
                    break
            if target is None:
                doc.add_heading(cat, level=2)
                doc.add_paragraph(text, style='List Bullet')
            else:
                # 插到该标题后面第一段之前的位置：简单起见追加到文末同类
                doc.add_paragraph('%s｜%s' % (cat, text), style='List Bullet')
            doc.save(path)
            return True
        except Exception:
            return False

    @filter.on_llm_response()
    async def _learn_phrase(self, event, resp) -> None:
        try:
            text = getattr(resp, 'completion_text', None)
            if text is None:
                text = str(resp)
            text = str(text).strip()
            if not (4 <= len(text) <= 60):
                return
            kw = {'难过': ['难过', '委屈', '算了', '不想'], '生气': ['随便', '哦', '不管'],
                  '开心': ['挺好', '喜欢', '不错', '好耶'], '困': ['困', '睡', '晚安'],
                  '害羞': ['别', '不好意思', '羞'], '关心': ['记得', '早点', '小心'],
                  '撒娇': ['陪我', '别走', '抱']}
            cat = '她说过的'
            for c, words in kw.items():
                if any(w in text for w in words):
                    cat = c
                    break
            self._corpus_add(cat, text)
        except Exception as e:
            logger.warning('learn_phrase failed: %s' % e)
