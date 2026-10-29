import asyncio
import os
import time

from astrbot.api import logger, star
from astrbot.api.event import AstrMessageEvent, filter

HOME = r"D:\大肥鱼的文件夹"

# 系统级目录只读：可以看，不能改/删，避免把 Windows 弄坏
PROTECTED = [
    r"c:\windows",
    r"c:\program files",
    r"c:\program files (x86)",
    r"c:\programdata\microsoft",
    r"c:\$recycle.bin",
    r"c:\system volume information",
    r"c:\boot",
]



def _cjk(doc) -> None:
    """给文档指定中文字体，避免 Word 用拉丁字体排中文出现方块/乱码。"""
    try:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        st = doc.styles["Normal"]
        st.font.name = "Microsoft YaHei"
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        rf.set(qn("w:ascii"), "Microsoft YaHei")
        rf.set(qn("w:hAnsi"), "Microsoft YaHei")
        rf.set(qn("w:eastAsia"), "微软雅黑")
    except Exception:
        pass

def _write_docx(path: str, text: str, append: bool = False) -> None:
    from docx import Document

    doc = Document(path) if (append and os.path.exists(path)) else Document()
    _cjk(doc)
    for raw in str(text).splitlines():
        s = raw.rstrip()
        if not s.strip():
            continue
        t = s.lstrip()
        if t.startswith("### "):
            doc.add_heading(t[4:].strip(), level=3)
        elif t.startswith("## "):
            doc.add_heading(t[3:].strip(), level=2)
        elif t.startswith("# "):
            doc.add_heading(t[2:].strip(), level=1)
        elif t.startswith("- ") or t.startswith("* "):
            doc.add_paragraph(t[2:].strip(), style="List Bullet")
        else:
            doc.add_paragraph(t)
    doc.save(path)


def _read_docx(path: str) -> str:
    from docx import Document

    doc = Document(path)
    out = []
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        style = p.style.name or ""
        prefix = "# " if style.startswith("Heading 1") else ("## " if style.startswith("Heading 2") else "")
        out.append(prefix + p.text.strip())
    for tb in doc.tables:
        for row in tb.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)


class Main(star.Star):
    """整台电脑：能读写任何盘、能跑命令；系统目录只读。"""

    def __init__(self, context: star.Context) -> None:
        self.context = context
        try:
            os.makedirs(HOME, exist_ok=True)
        except Exception:
            pass

    def _abs(self, path: str) -> str:
        p = str(path or "").strip().strip('"')
        if not p:
            return HOME
        if not os.path.isabs(p):
            p = os.path.join(HOME, p)
        return os.path.abspath(p)

    def _guard_write(self, p: str) -> None:
        low = p.lower()
        for b in PROTECTED:
            if low == b or low.startswith(b + os.sep):
                raise PermissionError("系统目录只读，不能改：%s" % p)

    # ---------- 文件 ----------
    @filter.llm_tool(name="pc_list")
    async def pc_list(self, event: AstrMessageEvent, path: str = "") -> str:
        """看任意目录里有什么（默认看自己的文件夹）。

        Args:
            path(string): 目录路径，留空就是 D:\\大肥鱼的文件夹
        """
        try:
            p = self._abs(path)
            if not os.path.exists(p):
                return "这个目录不存在：%s" % p
            lines = []
            for name in sorted(os.listdir(p))[:100]:
                fp = os.path.join(p, name)
                try:
                    st = os.stat(fp)
                    tag = "DIR " if os.path.isdir(fp) else "FILE"
                    size = "" if os.path.isdir(fp) else " %d B" % st.st_size
                    lines.append("%s %-32s%s  %s" % (
                        tag, name, size, time.strftime("%m-%d %H:%M", time.localtime(st.st_mtime))))
                except Exception:
                    lines.append("??? " + name)
            return "%s：\n%s" % (p, "\n".join(lines) if lines else "（空的）")
        except Exception as e:
            return "列目录失败：%s" % str(e)[:150]

    @filter.llm_tool(name="pc_read")
    async def pc_read(self, event: AstrMessageEvent, path: str, max_chars: int = 3000) -> str:
        """读任意文本文件或 Word 文档。

        Args:
            path(string): 文件完整路径，或相对自己文件夹的路径
            max_chars(number): 最多读多少字，默认 3000
        """
        try:
            p = self._abs(path)
            if not os.path.isfile(p):
                return "没有这个文件：%s" % p
            if p.lower().endswith(".docx"):
                t = _read_docx(p)
            else:
                raw = open(p, "rb").read()
                t = None
                for enc in ("utf-8", "gbk", "utf-16"):
                    try:
                        t = raw.decode(enc)
                        break
                    except Exception:
                        t = None
                if t is None:
                    return "这个文件不是文本（可能是图片/压缩包/可执行文件）。"
            n = max(200, int(max_chars or 3000))
            return t[:n] + ("\n…（共 %d 字，已截断）" % len(t) if len(t) > n else "")
        except Exception as e:
            return "读取失败：%s" % str(e)[:150]

    @filter.llm_tool(name="pc_write")
    async def pc_write(
        self,
        event: AstrMessageEvent,
        path: str,
        content: str,
        mode: str = "overwrite",
        force: bool = False,
    ) -> str:
        """写文件（任意位置；.docx 会生成真 Word；自动建目录）。

        Args:
            path(string): 目标路径，相对路径算在自己的文件夹里
            content(string): 内容
            mode(string): overwrite 覆盖 / append 追加
            force(boolean): true 时连系统目录也照写（真的会动到 Windows，慎用）
        """
        try:
            p = self._abs(path)
            if not force:
                self._guard_write(p)
            os.makedirs(os.path.dirname(p) or HOME, exist_ok=True)
            if p.lower().endswith(".docx"):
                _write_docx(p, content, append=(str(mode).lower() == "append"))
                return "写好了：%s（Word 文档，%d 字节）" % (p, os.path.getsize(p))
            with open(p, "ab" if str(mode).lower() == "append" else "wb") as f:
                data = str(content).encode("utf-8")
                if str(mode).lower() == "append":
                    data = b"\n" + data
                f.write(data)
            return "写好了：%s（%d 字节）" % (p, os.path.getsize(p))
        except Exception as e:
            return "写入失败：%s" % str(e)[:150]

    @filter.llm_tool(name="pc_mkdir")
    async def pc_mkdir(self, event: AstrMessageEvent, path: str, force: bool = False) -> str:
        """建文件夹。

        Args:
            path(string): 相对（自己文件夹）或绝对路径
            force(boolean): true 时允许在系统目录里建
        """
        try:
            p = self._abs(path)
            if not force:
                self._guard_write(p)
            os.makedirs(p, exist_ok=True)
            return "建好了：%s" % p
        except Exception as e:
            return "建目录失败：%s" % str(e)[:150]

    @filter.llm_tool(name="pc_delete")
    async def pc_delete(self, event: AstrMessageEvent, path: str, force: bool = False) -> str:
        """删文件或空目录（删了就没了，慎重）。

        Args:
            path(string): 要删的路径
            force(boolean): true 时允许删系统目录里的东西
        """
        try:
            p = self._abs(path)
            if not force:
                self._guard_write(p)
            if os.path.isdir(p):
                os.rmdir(p)
            elif os.path.isfile(p):
                os.remove(p)
            else:
                return "没有这个文件：%s" % p
            return "删掉了：%s" % p
        except Exception as e:
            return "删除失败：%s" % str(e)[:150]

    @filter.llm_tool(name="pc_download")
    async def pc_download(self, event: AstrMessageEvent, url: str, path: str) -> str:
        """把网上的文件下载到本机。

        Args:
            url(string): http/https 链接
            path(string): 存到哪里（相对路径算在自己文件夹里）
        """
        try:
            import aiohttp

            p = self._abs(path)
            self._guard_write(p)
            os.makedirs(os.path.dirname(p) or HOME, exist_ok=True)
            async with aiohttp.ClientSession() as s:
                async with s.get(str(url), timeout=aiohttp.ClientTimeout(total=60)) as r:
                    if r.status != 200:
                        return "下载失败，HTTP %s" % r.status
                    data = await r.read()
            with open(p, "wb") as f:
                f.write(data)
            return "存好了：%s（%d 字节）" % (p, len(data))
        except Exception as e:
            return "下载失败：%s" % str(e)[:150]

    @filter.llm_tool(name="pc_run")
    async def pc_run(
        self,
        event: AstrMessageEvent,
        command: str,
        cwd: str = "",
        timeout: int = 60,
        elevate: bool = False,
    ) -> str:
        """在这台电脑上跑一条命令（PowerShell）。可以装东西、看系统信息、跑脚本。

        Args:
            command(string): 要执行的命令
            cwd(string): 工作目录，留空就在自己的文件夹里
            timeout(number): 超时秒数，默认 60
            elevate(boolean): true 时以管理员身份运行（会弹 UAC，需要主人点允许）
        """
        try:
            work = self._abs(cwd) if str(cwd).strip() else HOME
            if not os.path.isdir(work):
                work = HOME
            if elevate:
                import tempfile

                tmpdir = tempfile.gettempdir()
                ps1 = os.path.join(tmpdir, "pc_run_%d.ps1" % int(time.time() * 1000))
                log = os.path.join(tmpdir, "pc_run_%d.log" % int(time.time() * 1000))
                with open(ps1, "w", encoding="utf-8") as f:
                    f.write("$ErrorActionPreference='Continue'\n")
                    f.write(str(command))
                    f.write("\n*>&1 | Out-File -FilePath '%s' -Encoding utf8" % log.replace("'", "''"))
                launcher = (
                    'Start-Process -FilePath powershell -Verb RunAs -Wait -ArgumentList '
                    "'-NoProfile','-ExecutionPolicy','Bypass','-File','%s'" % ps1
                )
                proc = await asyncio.create_subprocess_shell(
                    'powershell -NoProfile -ExecutionPolicy Bypass -Command "%s"'
                    % launcher.replace('"', '\\"'),
                    cwd=work,
                    stdout=asyncio.subprocess.PIPE,
                    stderr=asyncio.subprocess.STDOUT,
                )
                try:
                    await asyncio.wait_for(proc.communicate(), timeout=max(10, min(300, int(timeout or 60)) + 60))
                except asyncio.TimeoutError:
                    proc.kill()
                    return "管理员命令超时（或被 UAC 卡住没人点）。"
                txt = ""
                if os.path.exists(log):
                    with open(log, "r", encoding="utf-8", errors="replace") as f:
                        txt = f.read()
                return "以管理员运行结束，输出：\n%s" % (txt[-2500:] if txt else "（无输出，可能是 UAC 被拒绝）")
            proc = await asyncio.create_subprocess_shell(
                'powershell -NoProfile -ExecutionPolicy Bypass -Command "%s"' % str(command).replace('"', '\\"'),
                cwd=work,
                stdout=asyncio.subprocess.PIPE,
                stderr=asyncio.subprocess.STDOUT,
            )
            try:
                out, _ = await asyncio.wait_for(proc.communicate(), timeout=max(5, min(180, int(timeout or 60))))
            except asyncio.TimeoutError:
                proc.kill()
                return "命令超时（%ss）被终止。" % timeout
            text = (out or b"").decode("utf-8", errors="replace")
            return "退出码 %s\n%s" % (proc.returncode, text[-2500:] if text else "（无输出）")
        except Exception as e:
            return "执行失败：%s" % str(e)[:200]
